"""Excel availability upload with conflict detection."""

from __future__ import annotations

import io
import logging
from datetime import date, timedelta
from typing import Any, Dict, List, Optional

from app.reservations import store

logger = logging.getLogger(__name__)

EXPECTED_HEADERS = (
    "guesthouse_name",
    "room_number",
    "date_from",
    "date_until",
    "status",
)


def _parse_date(value: Any) -> str:
    if value is None:
        raise ValueError("Missing date")
    if hasattr(value, "date"):
        return value.date().isoformat()
    text = str(value).strip()
    if "T" in text:
        text = text.split("T")[0]
    return date.fromisoformat(text).isoformat()


def _date_range(from_s: str, until_s: str) -> List[str]:
    start = date.fromisoformat(from_s)
    end = date.fromisoformat(until_s)
    if end < start:
        raise ValueError("date_until must be on or after date_from")
    days: List[str] = []
    current = start
    while current <= end:
        days.append(current.isoformat())
        current += timedelta(days=1)
    return days


def parse_excel_rows(file_bytes: bytes) -> List[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel upload") from exc

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [str(h or "").strip().lower().replace(" ", "_") for h in header_row]
    idx = {h: i for i, h in enumerate(headers) if h}
    missing = [h for h in EXPECTED_HEADERS if h not in idx]
    if missing:
        raise ValueError(f"Missing columns: {', '.join(missing)}")

    parsed: List[dict] = []
    for row in rows_iter:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        parsed.append(
            {
                "guesthouse_name": str(row[idx["guesthouse_name"]] or "").strip(),
                "room_number": str(row[idx["room_number"]] or "").strip(),
                "date_from": _parse_date(row[idx["date_from"]]),
                "date_until": _parse_date(row[idx["date_until"]]),
                "status": str(row[idx["status"]] or "available").strip().lower(),
            }
        )
    return parsed


def _lookup_room(guesthouse_name: str, room_number: str) -> Optional[dict]:
    gh = store.get_guesthouse_by_name(guesthouse_name)
    if not gh:
        return None
    return store.get_room_by_guesthouse_number(gh["id"], room_number)


def process_excel_upload(
    file_bytes: bytes,
    uploaded_by: str,
) -> dict:
    """
    Parse Excel and apply availability where safe.
    Returns { applied, conflicts }.
    """
    rows = parse_excel_rows(file_bytes)
    applied = 0
    conflicts: List[dict] = []

    for row in rows:
        room = _lookup_room(row["guesthouse_name"], row["room_number"])
        if not room:
            conflicts.append(
                {
                    **row,
                    "reason": "Unknown guesthouse or room",
                }
            )
            continue

        status = row["status"]
        if status not in (
            store.AVAIL_AVAILABLE,
            store.AVAIL_BOOKED,
            store.AVAIL_MAINTENANCE,
            store.AVAIL_BLOCKED,
        ):
            conflicts.append({**row, "reason": f"Invalid status: {status}"})
            continue

        for day in _date_range(row["date_from"], row["date_until"]):
            conflict = _check_and_apply_day(
                room=room,
                day=day,
                status=status,
                uploaded_by=uploaded_by,
            )
            if conflict:
                conflicts.append(conflict)
            else:
                applied += 1

    return {"applied": applied, "conflicts": conflicts}


def _check_and_apply_day(
    *,
    room: dict,
    day: str,
    status: str,
    uploaded_by: str,
) -> Optional[dict]:
    calendar = store.get_room_availability_calendar(room["id"], day, day)
    current = calendar[0] if calendar else {"status": store.AVAIL_AVAILABLE}
    cur_status = current.get("status") or store.AVAIL_AVAILABLE
    reservation_id = current.get("reservation_id")

    if cur_status == store.AVAIL_BOOKED and reservation_id:
        res = store.get_reservation(reservation_id)
        if res and res.get("status") in (store.STATUS_CONFIRMED, store.STATUS_PENDING):
            return {
                "room_id": room["id"],
                "guesthouse_name": room.get("guesthouse_name"),
                "room_number": room.get("room_number"),
                "date": day,
                "proposed_status": status,
                "employee_email": res.get("employee_email"),
                "confirmation_number": res.get("confirmation_number"),
                "reservation_id": reservation_id,
                "reservation_status": res.get("status"),
            }

    store.apply_availability_status(
        room["id"],
        day,
        status,
        uploaded_by=uploaded_by,
        source="excel_upload",
    )
    return None


def resolve_conflicts(
    resolutions: List[dict],
    resolved_by: str,
) -> dict:
    """
    Each resolution: { room_id, date, proposed_status, action: skip|override, reservation_id?, reason? }
    """
    applied = 0
    skipped = 0
    for item in resolutions:
        action = (item.get("action") or "skip").lower()
        if action == "skip":
            skipped += 1
            continue
        if action == "override":
            rid = item.get("reservation_id")
            reason = item.get("reason") or "HR override due to availability upload"
            if rid:
                store.override_reservation(rid, resolved_by, reason)
            store.apply_availability_status(
                item["room_id"],
                item["date"],
                item.get("proposed_status") or store.AVAIL_MAINTENANCE,
                uploaded_by=resolved_by,
                source="excel_upload",
            )
            applied += 1
    return {"applied": applied, "skipped": skipped}
