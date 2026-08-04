"""Extract plain text from PDF / DOCX / TXT uploads."""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt"})


class UnsupportedDocumentType(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


def _ext(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise RuntimeError("PyMuPDF (pymupdf) is required for PDF extraction") from exc

    parts: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            parts.append(page.get_text("text") or "")
    return "\n".join(parts)


def extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX extraction") from exc

    from io import BytesIO

    document = Document(BytesIO(data))
    paras = [p.text for p in document.paragraphs if (p.text or "").strip()]
    # Tables
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    return "\n".join(paras)


def extract_text(data: bytes, filename: str) -> str:
    """Return raw extracted text for a supported file type."""
    if not data:
        raise EmptyDocumentError("Uploaded file is empty")
    ext = _ext(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentType(
            f"Unsupported file type '{ext or '(none)'}'. "
            f"Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if ext == ".pdf":
        raw = extract_pdf(data)
    elif ext == ".docx":
        raw = extract_docx(data)
    else:
        raw = extract_txt(data)

    if not (raw or "").strip():
        raise EmptyDocumentError("No extractable text found in the document")
    return raw
